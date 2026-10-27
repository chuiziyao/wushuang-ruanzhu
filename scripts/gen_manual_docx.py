# -*- coding: utf-8 -*-
"""把 Markdown 版操作说明书转成软著「文档鉴别材料」Word。

要点：
  · 页眉左＝软件全称 + 版本号，右＝第 X 页 共 N 页（N 用 NUMPAGES 域自动算）
  · 每章（## ）可选另起一页，便于按章补截图
  · 图片不存在时留占位文字，提醒补真实界面截图（不得用 AI 生成图）
  · 文档页数由 Word 排版决定，若超过 60 页，提交时按「前 30 页 + 后 30 页」截取 PDF 页码区间

用法：
  python gen_manual_docx.py --md 操作说明书.md --name "示例事项软件" --version V1.0 \
      --out 操作说明书.docx [--chapter-pagebreak]
"""
import argparse
import os
import re
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import config as cfgmod  # noqa: E402
from docx_common import (  # noqa: E402
    new_document, setup_page, build_header, set_run_font, set_update_fields,
    usable_width_cm,
)

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

BODY_PT = 10.5
MONO_PT = 9.0

INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_inline(par, text, size=BODY_PT, cn='宋体', mono='Consolas', base_bold=False):
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**') and len(seg) > 4:
            run = par.add_run(seg[2:-2])
            set_run_font(run, cn, cn, size, bold=True)
        elif seg.startswith('`') and seg.endswith('`') and len(seg) > 2:
            run = par.add_run(seg[1:-1])
            set_run_font(run, mono, cn, size - 0.5)
        else:
            run = par.add_run(seg)
            set_run_font(run, cn, cn, size, bold=base_bold)


def body_par(doc, text, first_line_indent=True, size=BODY_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.35
    pf.alignment = align
    if first_line_indent:
        pf.first_line_indent = Pt(size * 2)
    add_inline(par, text, size)
    return par


def heading(doc, text, level):
    sizes = {1: 16, 2: 13.5, 3: 11.5, 4: 10.5}
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(12 if level <= 2 else 8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    pf.keep_with_next = True
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = par.add_run(text)
    set_run_font(run, '黑体', '黑体', sizes.get(level, 10.5), bold=(level <= 3))
    return par


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    width = usable_width_cm(doc.sections[0])
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            text = row[j] if j < len(row) else ''
            add_inline(par, text, 9.5)
            if i == 0:
                for run in par.runs:
                    run.font.bold = True
            cell.width = Cm(width / cols)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, path, caption, base_dir):
    full = path if os.path.isabs(path) else os.path.join(base_dir, path)
    width = usable_width_cm(doc.sections[0]) - 1.0
    if os.path.isfile(full):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(full, width=Cm(width))
    else:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(2)
        run = par.add_run('【此处插入界面截图：%s】' % (caption or os.path.basename(path)))
        set_run_font(run, '宋体', '宋体', 10)
        run.font.italic = True
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        run = cap.add_run(caption)
        set_run_font(run, '宋体', '宋体', 9)


def add_code_block(doc, lines):
    for line in lines:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(13)
        pf.left_indent = Pt(12)
        run = par.add_run(line if line.strip() else ' ')
        set_run_font(run, 'Consolas', '宋体', MONO_PT)


def convert(md_path, out_path, name, version, chapter_pagebreak=False):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = new_document('宋体', '宋体', BODY_PT)
    section = doc.sections[0]
    setup_page(section, top_cm=2.2, bottom_cm=1.8, left_cm=2.2, right_cm=2.2)
    build_header(section, '%s %s' % (name.strip(), version.strip()), total_pages=None,
                 font_ascii='宋体', font_cn='宋体')
    set_update_fields(doc)
    base_dir = os.path.dirname(os.path.abspath(md_path))

    i, n = 0, len(lines)
    first_h1_done = False
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith('```'):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith('```'):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            add_code_block(doc, block)
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            par = heading(doc, text, level)
            if level == 1 and not first_h1_done:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_h1_done = True
            if level == 2 and chapter_pagebreak and first_h1_done:
                par.paragraph_format.page_break_before = True
            i += 1
            continue

        # 图片：![说明](路径)
        m = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if m:
            add_image(doc, m.group(2), m.group(1), base_dir)
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and stripped.count('|') >= 2:
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                row = lines[i].strip().strip('|')
                cells = [c.strip() for c in row.split('|')]
                if not all(re.match(r'^:?-{2,}:?$', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            par = body_par(doc, '· ' + m.group(1), first_line_indent=False)
            par.paragraph_format.left_indent = Pt(18)
            i += 1
            continue

        m = re.match(r'^(\d+)[.、)]\s+(.*)$', stripped)
        if m:
            par = body_par(doc, '%s. %s' % (m.group(1), m.group(2)), first_line_indent=False)
            par.paragraph_format.left_indent = Pt(18)
            i += 1
            continue

        if stripped.startswith('>'):
            par = body_par(doc, stripped.lstrip('> '), first_line_indent=False)
            par.paragraph_format.left_indent = Pt(18)
            for run in par.runs:
                run.font.italic = True
            i += 1
            continue

        body_par(doc, stripped)
        i += 1

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description='生成软著文档鉴别材料 Word（操作说明书）')
    ap.add_argument('--md', default=None, help='Markdown 说明书（缺省取当前目录 操作说明书.md）')
    ap.add_argument('--name', default=None, help='软件全称（缺省取 00_基本信息.json）')
    ap.add_argument('--version', default=None, help='版本号（缺省取配置，默认 V1.0）')
    ap.add_argument('--out', default=None, help='输出 docx 路径')
    ap.add_argument('--config', default=None, help='配置文件路径 00_基本信息.json')
    ap.add_argument('--chapter-pagebreak', action='store_true', help='每章另起一页')
    args = ap.parse_args()

    cfg = cfgmod.load(args.config)
    args.name = args.name or cfgmod.get(cfg, '软件全称')
    args.version = args.version or cfgmod.get(cfg, '版本号', 'V1.0')
    args.md = args.md or ('操作说明书.md' if os.path.isfile('操作说明书.md') else None)
    if not args.md:
        sys.exit('缺少说明书 Markdown：加 --md 操作说明书.md，或用 examples/操作说明书模板.md 改写')
    if not args.name:
        sys.exit('缺少软件全称：加 --name "全称"，或先 python bootstrap.py --name "全称" ...')

    out = args.out or os.path.splitext(args.md)[0] + '.docx'
    convert(args.md, out, args.name, args.version, args.chapter_pagebreak)
    print('已生成：%s' % out)
    print('提醒：文档页数由 Word 排版决定；若超过 60 页，提交时导出 PDF 后截取'
          '第 1–30 页与最后 30 页（页码保持原文真实页码）。')


if __name__ == '__main__':
    main()
